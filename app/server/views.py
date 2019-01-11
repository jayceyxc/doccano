import csv
import json
import logging
import xlrd
import xlwt
import openpyxl
from io import TextIOWrapper

from django.urls import reverse
from django.http import HttpResponse, HttpResponseRedirect
from django.shortcuts import get_object_or_404
from django.views import View
from django.views.generic import TemplateView, CreateView
from django.views.generic.list import ListView
from django.contrib.auth.mixins import LoginRequiredMixin

from .permissions import SuperUserMixin
from .forms import ProjectForm
from .models import Document, Project


class IndexView(TemplateView):
    template_name = 'index.html'


class ProjectView(LoginRequiredMixin, TemplateView):

    def get_template_names(self):
        project = get_object_or_404(Project, pk=self.kwargs['project_id'])
        return [project.get_template_name()]


class ProjectsView(LoginRequiredMixin, CreateView):
    form_class = ProjectForm
    template_name = 'projects.html'


class DatasetView(SuperUserMixin, LoginRequiredMixin, ListView):
    template_name = 'admin/dataset.html'
    paginate_by = 5

    def get_queryset(self):
        project = get_object_or_404(Project, pk=self.kwargs['project_id'])
        return project.documents.all()


class LabelView(SuperUserMixin, LoginRequiredMixin, TemplateView):
    template_name = 'admin/label.html'


class StatsView(SuperUserMixin, LoginRequiredMixin, TemplateView):
    template_name = 'admin/stats.html'


class GuidelineView(SuperUserMixin, LoginRequiredMixin, TemplateView):
    template_name = 'admin/guideline.html'


class DataUpload(SuperUserMixin, LoginRequiredMixin, TemplateView):
    template_name = 'admin/dataset_upload.html'

    def __init__(self):
        super(DataUpload, self).__init__()
        self.logger = logging.getLogger('log')

    def post(self, request, *args, **kwargs):
        project = get_object_or_404(Project, pk=kwargs.get('project_id'))
        import_format = request.POST['format']
        try:
            self.logger.info('upload format: ' + import_format)
            upload_file_name = request.FILES['file'].name
            self.logger.info(u'上传文件名：' + upload_file_name)
            if import_format == 'csv':
                form_data = TextIOWrapper(
                    request.FILES['file'].file, encoding='utf-8')
                reader = csv.reader(form_data)
                Document.objects.bulk_create([
                    Document(text=line[0].strip(), project=project)
                    for line in reader
                ])
            elif import_format == 'json':
                form_data = request.FILES['file'].file
                Document.objects.bulk_create([
                    Document(text=json.loads(entry)['text'], project=project)
                    for entry in form_data
                ])
            elif import_format == 'txt':
                form_data = TextIOWrapper(
                    request.FILES['file'].file, encoding='utf-8')
                Document.objects.bulk_create([
                    Document(text=line.strip(), project=project)
                    for line in form_data
                ])
            elif import_format == 'excel':
                form_data = TextIOWrapper(
                    request.FILES['file'].file, encoding='utf-8')
                data = form_data.readlines()
                workbook = xlrd.open_workbook(file_contents=form_data.readlines())
                # openpyxl.load_workbook(filename=form_data)
                for sheet_name in workbook.sheet_names():
                    worksheet = workbook.sheet_by_name(sheet_name=sheet_name)
                    for i in range(worksheet.nrows):
                        content = worksheet.cell(colx=0, rowx=i).value
                        content = content.strip()
                        self.logger.info(content)

            return HttpResponseRedirect(reverse('dataset', args=[project.id]))
        except Exception as e:
            self.logger.error('upload error', e.__cause__)
            return HttpResponseRedirect(reverse('upload', args=[project.id]))


class DataDownload(SuperUserMixin, LoginRequiredMixin, TemplateView):
    template_name = 'admin/dataset_download.html'


class DataDownloadFile(SuperUserMixin, LoginRequiredMixin, View):

    def __init__(self):
        super(DataDownloadFile, self).__init__()
        self.logger = logging.getLogger('log')

    def get(self, request, *args, **kwargs):
        project_id = self.kwargs['project_id']
        project = get_object_or_404(Project, pk=project_id)
        docs = project.get_documents(is_null=False).distinct()
        export_format = request.GET.get('format')
        filename = '_'.join(project.name.lower().split())
        self.logger.info('download data file format: ' + export_format)
        try:
            if export_format == 'csv':
                response = self.get_csv(filename, docs)
            elif export_format == 'json':
                response = self.get_json(filename, docs)
            elif export_format == 'bio':
                response = self.get_bio_text(filename, docs)
            return response
        except:
            return HttpResponseRedirect(reverse('download', args=[project.id]))

    def get_csv(self, filename, docs):
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="{}.csv"'.format(filename)
        writer = csv.writer(response)
        for d in docs:
            writer.writerows(d.to_csv())
        return response

    def get_json(self, filename, docs):
        response = HttpResponse(content_type='text/json')
        response['Content-Disposition'] = 'attachment; filename="{}.json"'.format(filename)
        for d in docs:
            dump = json.dumps(d.to_json(), ensure_ascii=False)
            response.write(dump + '\n') # write each json object end with a newline
        print('dump done')
        return response

    def get_bio_text(self, filename, docs):
        response = HttpResponse(content_type='text/plain')
        response['Content-Disposition'] = 'attachment; filename="{}.txt"'.format(filename)
        for d in docs:
            response.write(d.to_bio() + '\n') # write each text end with a newline
        print('dump done')
        return response


class DemoTextClassification(TemplateView):
    template_name = 'demo/demo_text_classification.html'


class DemoNamedEntityRecognition(TemplateView):
    template_name = 'demo/demo_named_entity.html'


class DemoTranslation(TemplateView):
    template_name = 'demo/demo_translation.html'
